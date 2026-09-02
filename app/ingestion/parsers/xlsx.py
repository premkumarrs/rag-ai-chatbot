"""XLSX parser."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.config import CSV_ROWS_PER_CHUNK
from app.ingestion.parsers.base import BaseParser, build_parsed_document


class XLSXParser(BaseParser):
    file_type = "xlsx"
    extensions = {".xlsx"}

    def parse(self, file_path: Path, source_path: str):
        workbook = load_workbook(filename=str(file_path), read_only=True, data_only=True)
        sections: list[tuple[str, dict[str, object]]] = []

        for sheet in workbook.worksheets:
            rows = [
                ["" if cell is None else str(cell) for cell in row]
                for row in sheet.iter_rows(values_only=True)
            ]
            rows = [row for row in rows if any(str(value).strip() for value in row)]
            if not rows:
                continue

            headers = [str(value).strip() for value in rows[0]]
            data_rows = rows[1:] if len(rows) > 1 else []
            header_line = " | ".join(headers)

            for start in range(0, len(data_rows), CSV_ROWS_PER_CHUNK):
                batch = data_rows[start : start + CSV_ROWS_PER_CHUNK]
                lines = [header_line]
                for row in batch:
                    pairs = [
                        f"{headers[index]}: {row[index]}"
                        for index in range(min(len(headers), len(row)))
                        if str(row[index]).strip()
                    ]
                    if pairs:
                        lines.append(" | ".join(pairs))
                sections.append(
                    (
                        "\n".join(lines),
                        {
                            "source_path": source_path,
                            "filename": file_path.name,
                            "file_type": self.file_type,
                            "sheet_name": sheet.title,
                            "row_start": start + 1,
                            "row_end": start + len(batch),
                        },
                    )
                )

        workbook.close()
        return build_parsed_document(file_path, source_path, self.file_type, sections)
