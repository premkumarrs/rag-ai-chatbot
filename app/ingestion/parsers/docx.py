"""DOCX parser."""

from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument
from docx.text.paragraph import Paragraph

from app.ingestion.parsers.base import BaseParser, build_parsed_document


def _paragraph_heading(paragraph: Paragraph) -> str | None:
    style_name = paragraph.style.name if paragraph.style else ""
    if style_name.startswith("Heading"):
        return paragraph.text.strip()
    return None


class DOCXParser(BaseParser):
    file_type = "docx"
    extensions = {".docx"}

    def parse(self, file_path: Path, source_path: str):
        document = DocxDocument(str(file_path))
        sections: list[tuple[str, dict[str, object]]] = []
        current_heading: str | None = None
        paragraph_parts: list[str] = []

        def flush_paragraphs() -> None:
            nonlocal paragraph_parts
            if not paragraph_parts:
                return
            sections.append(
                (
                    "\n".join(paragraph_parts),
                    {
                        "source_path": source_path,
                        "filename": file_path.name,
                        "file_type": self.file_type,
                        "section": current_heading,
                    },
                )
            )
            paragraph_parts = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            heading = _paragraph_heading(paragraph)
            if heading:
                flush_paragraphs()
                current_heading = heading
                sections.append(
                    (
                        text,
                        {
                            "source_path": source_path,
                            "filename": file_path.name,
                            "file_type": self.file_type,
                            "section": heading,
                            "content_type": "heading",
                        },
                    )
                )
                continue
            paragraph_parts.append(text)

        flush_paragraphs()

        for table_index, table in enumerate(document.tables, start=1):
            rows: list[str] = []
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip() for cell in row.cells if cell.text.strip()
                )
                if row_text:
                    rows.append(row_text)
            if rows:
                sections.append(
                    (
                        "\n".join(rows),
                        {
                            "source_path": source_path,
                            "filename": file_path.name,
                            "file_type": self.file_type,
                            "section": current_heading,
                            "content_type": "table",
                            "table_index": table_index,
                        },
                    )
                )

        return build_parsed_document(file_path, source_path, self.file_type, sections)
