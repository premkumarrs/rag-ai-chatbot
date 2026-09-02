"""PPTX parser."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from app.ingestion.parsers.base import BaseParser, build_parsed_document


class PPTXParser(BaseParser):
    file_type = "pptx"
    extensions = {".pptx"}

    def parse(self, file_path: Path, source_path: str):
        presentation = Presentation(str(file_path))
        sections: list[tuple[str, dict[str, object]]] = []

        for slide_number, slide in enumerate(presentation.slides, start=1):
            title = ""
            bullets: list[str] = []
            tables: list[str] = []

            if slide.shapes.title and slide.shapes.title.text:
                title = slide.shapes.title.text.strip()

            for shape in slide.shapes:
                if not hasattr(shape, "text"):
                    continue
                text = shape.text.strip()
                if not text:
                    continue
                if title and text == title:
                    continue
                bullets.append(text)

                if shape.has_table:
                    for row in shape.table.rows:
                        row_text = " | ".join(
                            cell.text.strip() for cell in row.cells if cell.text.strip()
                        )
                        if row_text:
                            tables.append(row_text)

            content_parts = []
            if title:
                content_parts.append(f"Title: {title}")
            if bullets:
                content_parts.extend(bullets)
            if tables:
                content_parts.extend(tables)

            if not content_parts:
                continue

            sections.append(
                (
                    "\n".join(content_parts),
                    {
                        "source_path": source_path,
                        "filename": file_path.name,
                        "file_type": self.file_type,
                        "slide_number": slide_number,
                        "title": title or None,
                    },
                )
            )

        return build_parsed_document(file_path, source_path, self.file_type, sections)
