"""Tests for multi-format ingestion parsers and pipeline behavior."""

from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

from app.ingestion.chunking import chunk_parsed_document
from app.ingestion.loader import discover_files, is_supported_extension, parse_file
from app.ingestion.models import DocumentSection, ParsedDocument
from app.ingestion.parsers import SUPPORTED_EXTENSIONS
from app.ingestion.parsers.base import file_content_hash
from app.ingestion.parsers.json import JSONParser
from app.ingestion.parsers.ocr import OCREngine
from app.ingestion.pipeline import ingest_file, run_ingestion
from app.ingestion.report import IngestionReport

FIXTURES = Path(__file__).parent / "fixtures"


class ParserTests(unittest.TestCase):
    def test_supported_extensions_include_all_target_formats(self) -> None:
        expected = {
            ".pdf",
            ".docx",
            ".txt",
            ".md",
            ".csv",
            ".xlsx",
            ".pptx",
            ".html",
            ".htm",
            ".json",
            ".png",
            ".jpg",
            ".jpeg",
            ".tif",
            ".tiff",
        }
        self.assertTrue(expected.issubset(SUPPORTED_EXTENSIONS))

    def test_txt_parser(self) -> None:
        parsed = parse_file(FIXTURES / "sample.txt", FIXTURES.parent.parent)
        assert parsed is not None
        self.assertTrue(parsed.has_content)
        self.assertEqual(parsed.file_type, "txt")

    def test_markdown_parser(self) -> None:
        parsed = parse_file(FIXTURES / "sample.md", FIXTURES.parent.parent)
        assert parsed is not None
        self.assertTrue(any("markdown paragraph" in s.content for s in parsed.sections))

    def test_csv_parser(self) -> None:
        parsed = parse_file(FIXTURES / "sample.csv", FIXTURES.parent.parent)
        assert parsed is not None
        self.assertTrue(any("Widget A" in s.content for s in parsed.sections))

    def test_html_parser(self) -> None:
        parsed = parse_file(FIXTURES / "sample.html", FIXTURES.parent.parent)
        assert parsed is not None
        self.assertTrue(any("Support Guide" in s.content for s in parsed.sections))
        self.assertFalse(any("<script>" in s.content for s in parsed.sections))

    def test_json_parser(self) -> None:
        parsed = parse_file(FIXTURES / "sample.json", FIXTURES.parent.parent)
        assert parsed is not None
        self.assertIn("product.name: ABC", parsed.sections[0].content)

    def test_invalid_json_raises(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            bad = Path(tmp) / "bad.json"
            bad.write_text("{invalid", encoding="utf-8")
            parser = JSONParser()
            with self.assertRaises(json.JSONDecodeError):
                parser.parse(bad, "data/bad.json")

    def test_xlsx_parser(self) -> None:
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sample.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Products"
            sheet.append(["Product", "Code"])
            sheet.append(["Widget", "W1"])
            workbook.save(path)
            parsed = parse_file(path, Path(tmp))
            assert parsed is not None
            self.assertEqual(parsed.file_type, "xlsx")
            self.assertIn("Products", parsed.sections[0].metadata.get("sheet_name", ""))

    def test_pptx_parser(self) -> None:
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sample.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "Troubleshooting"
            slide.placeholders[1].text = "Check the power cable."
            presentation.save(path)
            parsed = parse_file(path, Path(tmp))
            assert parsed is not None
            self.assertEqual(parsed.sections[0].metadata.get("slide_number"), 1)

    def test_docx_parser(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sample.docx"
            document = Document()
            document.add_heading("Warranty", level=1)
            document.add_paragraph("The warranty period is 12 months.")
            table = document.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "Model"
            table.rows[0].cells[1].text = "X100"
            document.save(path)
            parsed = parse_file(path, Path(tmp))
            assert parsed is not None
            combined = "\n".join(section.content for section in parsed.sections)
            self.assertIn("warranty period is 12 months", combined.lower())

    def test_unsupported_extension_returns_none(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "program.exe"
            path.write_bytes(b"MZ")
            self.assertFalse(is_supported_extension(path.suffix))
            self.assertIsNone(parse_file(path, Path(tmp)))


class ChunkingTests(unittest.TestCase):
    def test_chunking_preserves_metadata(self) -> None:
        parsed = ParsedDocument(
            source_path="data/sample.txt",
            filename="sample.txt",
            file_type="txt",
            title="sample",
            content_hash="abc",
            sections=[
                DocumentSection(
                    content="Paragraph one. " * 100,
                    metadata={"section": "Intro"},
                )
            ],
        )
        chunks = chunk_parsed_document(parsed)
        self.assertGreater(len(chunks), 1)
        self.assertEqual(chunks[0].metadata["file_type"], "txt")


class IdempotencyTests(unittest.TestCase):
    @patch("app.ingestion.pipeline.embed_documents")
    @patch("app.ingestion.pipeline.replace_document_chunks")
    @patch("app.ingestion.pipeline.upsert_document")
    @patch("app.ingestion.pipeline.get_stored_content_hash")
    @patch("app.ingestion.pipeline.parse_file")
    def test_unchanged_file_skips_reembedding(
        self,
        mock_parse: MagicMock,
        mock_hash: MagicMock,
        mock_upsert: MagicMock,
        mock_replace: MagicMock,
        mock_embed: MagicMock,
    ) -> None:
        report = IngestionReport()
        conn = MagicMock()
        parsed = parse_file(FIXTURES / "sample.txt", FIXTURES.parent.parent)
        assert parsed is not None
        mock_hash.return_value = parsed.content_hash

        ingest_file(conn, FIXTURES / "sample.txt", report)

        mock_parse.assert_not_called()
        mock_embed.assert_not_called()
        self.assertEqual(report.files_unchanged, 1)


class OCRTests(unittest.TestCase):
    def test_ocr_engine_reports_unavailable_without_tesseract(self) -> None:
        engine = OCREngine()
        if not engine.available:
            with self.assertRaises(RuntimeError):
                engine.extract_text_from_image(FIXTURES / "sample.txt")


class DiscoveryTests(unittest.TestCase):
    def test_nested_directory_discovery(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            nested = root / "support"
            nested.mkdir()
            (nested / "faq.txt").write_text("Help text", encoding="utf-8")
            cache_dir = root / "__pycache__"
            cache_dir.mkdir()
            (cache_dir / "ignore.txt").write_text("ignore", encoding="utf-8")
            files = discover_files(root)
            self.assertEqual(len(files), 1)


class PartialFailureTests(unittest.TestCase):
    @patch("app.ingestion.pipeline.get_connection")
    def test_one_bad_file_does_not_stop_ingestion(self, mock_conn_factory: MagicMock) -> None:
        conn = MagicMock()
        mock_conn_factory.return_value.__enter__.return_value = conn

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "good.txt").write_text("hello", encoding="utf-8")
            (root / "bad.json").write_text("{invalid", encoding="utf-8")
            report = run_ingestion(root)

        self.assertEqual(report.files_processed, 1)
        self.assertEqual(report.files_failed, 1)


if __name__ == "__main__":
    unittest.main()
